"""Build the editable patent discussion draft and its monochrome figures."""

from __future__ import annotations

import argparse
import re
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "造口术后腹壁训练视频分析方法_发明专利申请初稿.md"
OUTPUT = ROOT / "造口术后腹壁训练视频分析方法_发明专利申请初稿.docx"
FONT_FILE = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
CHINESE_FONT = "Arial Unicode MS"


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_FILE, size)


def center_text(draw: ImageDraw.ImageDraw, xy: tuple[int, int], value: str, fnt: ImageFont.FreeTypeFont) -> None:
    box = draw.textbbox((0, 0), value, font=fnt)
    draw.text((xy[0] - (box[2] - box[0]) / 2, xy[1] - (box[3] - box[1]) / 2), value, fill="black", font=fnt)


def box(draw: ImageDraw.ImageDraw, x0: int, y0: int, x1: int, y1: int, label: str) -> None:
    draw.rounded_rectangle((x0, y0, x1, y1), radius=12, outline="black", width=4)
    center_text(draw, ((x0 + x1) // 2, (y0 + y1) // 2), label, font(31))


def arrow(draw: ImageDraw.ImageDraw, x0: int, y0: int, x1: int, y1: int) -> None:
    draw.line((x0, y0, x1, y1), fill="black", width=4)
    if abs(y1 - y0) > abs(x1 - x0):
        s = 1 if y1 > y0 else -1
        draw.polygon([(x1, y1), (x1 - 12, y1 - 18 * s), (x1 + 12, y1 - 18 * s)], fill="black")
    else:
        s = 1 if x1 > x0 else -1
        draw.polygon([(x1, y1), (x1 - 18 * s, y1 - 12), (x1 - 18 * s, y1 + 12)], fill="black")


def save_figure_1(path: Path) -> None:
    im = Image.new("RGB", (1600, 820), "white")
    d = ImageDraw.Draw(im)
    labels = [
        "S101  获取视频与造口区域参数",
        "S102  建立逐帧躯干坐标",
        "S103  排除遮挡并选择观测条带",
        "S104  静息呼吸和噪声校准",
        "S105  计算局部轮廓残差",
        "S106  周期识别与证据门控",
    ]
    for i, label in enumerate(labels):
        y = 35 + i * 127
        box(d, 250, y, 1350, y + 82, label)
        if i < len(labels) - 1:
            arrow(d, 800, y + 84, 800, y + 122)
    im.save(path)


def save_figure_2(path: Path) -> None:
    im = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(im)
    d.ellipse((650, 45, 950, 250), outline="black", width=6)
    d.line((610, 260, 480, 420, 520, 815), fill="black", width=7, joint="curve")
    d.line((990, 260, 1120, 420, 1080, 815), fill="black", width=7, joint="curve")
    d.line((610, 260, 990, 260), fill="black", width=7)
    d.line((520, 815, 1080, 815), fill="black", width=7)
    d.line((800, 285, 800, 760), fill="black", width=3)
    d.line((745, 710, 1040, 710), fill="black", width=3)
    arrow(d, 800, 570, 800, 360)
    arrow(d, 800, 710, 1010, 710)
    d.text((820, 390), "y", font=font(32), fill="black")
    d.text((1000, 730), "x", font=font(32), fill="black")
    d.rectangle((905, 445, 1130, 655), outline="black", width=5)
    for x in range(905, 1130, 22):
        d.line((x, 445, min(x + 210, 1130), 655), fill="gray", width=2)
    d.rounded_rectangle((470, 420, 605, 570), radius=10, outline="black", width=4)
    d.rounded_rectangle((485, 600, 615, 750), radius=10, outline="black", width=4)
    d.text((120, 460), "A1 可见条带", font=font(32), fill="black")
    d.text((120, 650), "A2 可见条带", font=font(32), fill="black")
    d.line((360, 478, 472, 478), fill="black", width=3)
    d.line((360, 665, 485, 665), fill="black", width=3)
    d.text((1150, 500), "Zt 遮挡排除区", font=font(32), fill="black")
    d.line((1130, 540, 1148, 540), fill="black", width=3)
    im.save(path)


def save_figure_3(path: Path) -> None:
    im = Image.new("RGB", (1600, 760), "white")
    d = ImageDraw.Draw(im)
    x0, x1, y0 = 120, 1510, 625
    arrow(d, x0, y0, x1, y0)
    arrow(d, x0, y0, x0, 125)
    d.text((1440, 640), "时间 t", font=font(28), fill="black")
    d.text((150, 55), "代理信号", font=font(28), fill="black")
    bounds = [120, 420, 740, 1050, 1390]
    labels = ["准备", "变化", "保持", "回位"]
    for j, label in enumerate(labels):
        cx = (bounds[j] + bounds[j + 1]) // 2
        center_text(d, (cx, 690), label, font(32))
        if j > 0:
            d.line((bounds[j], 155, bounds[j], 625), fill="gray", width=2)
    pts = [(120, 510), (240, 506), (350, 515), (420, 505), (490, 455), (570, 380),
           (650, 300), (740, 250), (830, 246), (950, 250), (1050, 255),
           (1130, 330), (1230, 440), (1320, 500), (1390, 510)]
    d.line(pts, fill="black", width=7, joint="curve")
    for y in (480, 540):
        d.line((120, y, 1390, y), fill="gray", width=2)
    d.text((1420, 480), "静息区间", font=font(26), fill="black")
    im.save(path)


def save_figure_4(path: Path) -> None:
    im = Image.new("RGB", (1600, 850), "white")
    d = ImageDraw.Draw(im)
    box(d, 450, 35, 1150, 130, "候选动作周期 C")
    box(d, 400, 225, 1200, 325, "可观测性与基线有效？")
    arrow(d, 800, 132, 800, 220)
    box(d, 75, 515, 650, 620, "不可判定＋原因提示")
    d.line((400, 275, 225, 275, 225, 512), fill="black", width=4)
    d.text((255, 280), "否", font=font(30), fill="black")
    box(d, 840, 420, 1510, 520, "四状态闭环且局部变化有效？")
    d.line((1200, 275, 1360, 275, 1360, 417), fill="black", width=4)
    d.text((1220, 280), "是", font=font(30), fill="black")
    box(d, 785, 660, 1260, 765, "完整动作周期")
    arrow(d, 1030, 525, 1030, 655)
    d.text((1050, 565), "是", font=font(30), fill="black")
    box(d, 75, 675, 650, 780, "不完整／姿态运动主导")
    d.line((840, 470, 710, 470, 710, 730, 655, 730), fill="black", width=4)
    d.text((725, 485), "否", font=font(30), fill="black")
    im.save(path)


def set_east_asian_font(style, name: str) -> None:
    style.font.name = name
    style.font.color.rgb = None
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), name)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def main(source: Path = SOURCE, output: Path = OUTPUT) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.86)
    sec.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    set_east_asian_font(normal, CHINESE_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(7)
    for style_name, size, before, after in [
        ("Title", 17, 0, 14),
        ("Heading 1", 14, 12, 8),
        ("Heading 2", 12, 10, 6),
        ("Heading 3", 11, 8, 5),
    ]:
        style = doc.styles[style_name]
        set_east_asian_font(style, CHINESE_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.style = normal
    footer.add_run("发明专利申请文件讨论稿 · 2026-10-30")

    current_section = ""
    for line in source.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(line[2:], style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_keep_with_next(p)
            subtitle = doc.add_paragraph("发明专利申请文件（技术交底讨论稿）")
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.paragraph_format.space_after = Pt(22)
            continue
        if line.startswith("## "):
            current_section = line[3:]
            if current_section in ("权利要求书", "说明书"):
                doc.add_page_break()
            p = doc.add_paragraph(current_section, style="Heading 1")
            set_keep_with_next(p)
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(line[4:], style="Heading 2")
            set_keep_with_next(p)
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(line[5:], style="Heading 3")
            set_keep_with_next(p)
            continue
        p = doc.add_paragraph(line)
        if current_section == "权利要求书" and re.match(r"^\d+\.", line):
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(10)
        elif current_section == "说明书":
            p.paragraph_format.first_line_indent = Inches(0.26)

    with tempfile.TemporaryDirectory(prefix="patent-figures-") as temp:
        paths = [Path(temp) / f"figure-{i}.png" for i in range(1, 5)]
        for make, path in zip((save_figure_1, save_figure_2, save_figure_3, save_figure_4), paths):
            make(path)
        doc.add_page_break()
        p = doc.add_paragraph("说明书附图", style="Heading 1")
        set_keep_with_next(p)
        captions = [
            "图1 视频处理流程图",
            "图2 躯干坐标系、局部遮挡排除区及观测条带示意图",
            "图3 局部轮廓变化的四状态动作周期示意图",
            "图4 动作证据门控与不可判定处理示意图",
        ]
        for i, (path, caption) in enumerate(zip(paths, captions)):
            if i == 2:
                doc.add_page_break()
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.add_run().add_picture(str(path), width=Inches(6.55))
            pic.paragraph_format.space_after = Pt(2)
            p = doc.add_paragraph(caption)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(16)
        doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    main(args.source, args.output)
